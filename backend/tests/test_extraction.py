import pytest
from docx import Document

from app.services.extraction.base import ExtractionQualityError
from app.services.extraction.docx_extractor import DocxExtractor
from app.services.extraction.factory import get_extractor
from app.services.extraction.txt_extractor import TxtExtractor


def test_txt_extractor(tmp_path):
    p = tmp_path / "lesson.txt"
    p.write_text("A train travels 240 km in 4 hours. Speed = distance / time.")
    result = TxtExtractor().extract(str(p))
    assert "240 km" in result.text


def test_txt_extractor_rejects_empty(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("   ")
    with pytest.raises(ExtractionQualityError):
        TxtExtractor().extract(str(p))


def test_docx_extractor(tmp_path):
    p = tmp_path / "lesson.docx"
    doc = Document()
    doc.add_paragraph("Average speed is distance divided by time.")
    doc.add_paragraph("A cyclist travels 45 km in 3 hours.")
    doc.save(str(p))

    result = DocxExtractor().extract(str(p))
    assert "45 km" in result.text


def test_factory_unknown_type():
    with pytest.raises(ValueError):
        get_extractor("exe")


def test_factory_returns_correct_extractor():
    assert isinstance(get_extractor("txt"), TxtExtractor)
