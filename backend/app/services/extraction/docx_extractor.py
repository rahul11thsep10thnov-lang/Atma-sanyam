from docx import Document

from app.services.extraction.base import DocumentExtractor, ExtractionResult


class DocxExtractor(DocumentExtractor):
    def extract(self, file_path: str) -> ExtractionResult:
        doc = Document(file_path)
        parts: list[str] = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
        self._check_quality(text)
        return ExtractionResult(text=text, page_count=None, warnings=[])
